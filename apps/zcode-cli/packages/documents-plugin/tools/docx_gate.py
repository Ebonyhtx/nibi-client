"""DOCX 质量门：交付前把文档逐条验一遍，验不到的要显式标「跳过」而不是沉默放过。

## 判据分层

**A 层（静态）**
1. `font-size` 字号下限（正文 ≥9pt 建议、任何位置 <7pt 硬失败）
2. `fonts` 中文字体必须有**中文字体**（只设 ascii/hAnsi 会让 Word 替换字体）+ 字体数量上限
3. `style-consistency` 同名样式必须同字号同字体、标题层级不许跳级、正文不许到处手工格式
4. `images` 图片纵横比是否被拉扁、显示宽度/有效 DPI/替代文字
5. `toc` 目录条目、页码齐全（页码是否**真的是对的**在引擎层对账）
6. `page-setup` 纸张、页边距、页脚页码域、正文首行缩进
7. `hygiene` 交付卫生（本机路径/用户名/凭据、核心属性）
8. `structure` 空文档、尾部空段、长表重复表头、正文里用制表位排版

**B 层（引擎）**
9. `render` 转 PDF → 逐页文本抓缺字形方框与目录域错误、逐页 PNG 落证据目录，
   并把**目录页码与渲染页码逐条对账**

引擎缺席时 B 层判据是 `skip` + 原因；`--require-engine` 可升级为硬失败。
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from nibi_kit import cli, engines, gates, hygiene  # noqa: E402
from nibi_kit.env import require_modules  # noqa: E402

require_modules(["docx", "PIL"], "docx_gate（python-docx + pillow）")

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

import docx_fonts  # noqa: E402
import docx_spec  # noqa: E402
import docx_toc  # noqa: E402

EXPECTED_GATES = ["font-size", "fonts", "style-consistency", "images", "toc",
                  "page-setup", "hygiene", "structure", "render"]

FONT_FLOOR_HARD = 7.0
FONT_FLOOR_SOFT = 9.0
LATIN_ONLY = ("calibri", "timesnewroman", "timesnewromanpsmt", "arial", "cambria",
              "segoeui", "helvetica", "georgia", "verdana", "couriernew", "consolas",
              "dejavusans", "liberationserif", "liberationsans")
CJK_HINT = re.compile(r"[\u3000-\u303f\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\uff00-\uffef]")
TOC_ERRORS = ("错误！未定义书签", "Error! Bookmark not defined", "错误!未定义书签")


def gate_font_size(document, report: gates.GateReport) -> None:
    gate = report.gate("font-size", "字号下限")
    for location, paragraph, run in docx_fonts.iter_runs(document):
        resolved = docx_fonts.run_fonts(run, paragraph, document)   # 逐 run：段内混排也要查
        size = resolved["size"]
        gate.checked += 1
        if size < FONT_FLOOR_HARD:
            gate.fail("font-too-small", f"字号 {size}pt 小于下限 {FONT_FLOOR_HARD}pt", location)
        elif size < FONT_FLOOR_SOFT and location not in ("页脚1", "页眉1"):
            gate.warn("font-small", f"字号 {size}pt 偏小（正文建议 ≥{FONT_FLOOR_SOFT}pt）", location)


def gate_fonts(document, report: gates.GateReport) -> None:
    gate = report.gate("fonts", "中文字体与字体数量")
    cjk_families: dict[str, int] = {}
    for location, paragraph, run in docx_fonts.iter_runs(document):
        resolved = docx_fonts.run_fonts(run, paragraph, document)
        has_cjk = bool(CJK_HINT.search(run.text))
        gate.checked += 1
        if has_cjk:
            family = str(resolved["cjk"] or "")
            cjk_families[family] = cjk_families.get(family, 0) + 1
            if family.replace(" ", "").lower() in LATIN_ONLY:
                gate.fail("font-not-cjk",
                          f"含中文的文字用了拉丁字体 eastAsia={family!r}（Word 会替换字体）",
                          location)
    if len(cjk_families) > 4:
        gate.warn("too-many-fonts",
                  f"文档用了 {len(cjk_families)} 种中文字体（{'、'.join(sorted(cjk_families))}）"
                  f"—— 同级文本应统一字体")
    gate.notes.append("中文字体分布：" + "、".join(
        f"{name}×{count}" for name, count in sorted(cjk_families.items())) or "无中文内容")


def gate_style_consistency(document, report: gates.GateReport) -> None:
    gate = report.gate("style-consistency", "样式一致性")
    by_style: dict[str, list[tuple[str, float, str]]] = {}
    levels_seen: list[tuple[int, str]] = []
    direct = 0
    for location, paragraph, run in docx_fonts.iter_runs(document):
        resolved = docx_fonts.run_fonts(run, paragraph, document)
        name = resolved["styleName"] or "Normal"
        by_style.setdefault(name, []).append((location, resolved["size"], str(resolved["cjk"])))
        if resolved["directFormat"] and name in ("body", "bodyNoIndent"):
            direct += 1
        head = re.fullmatch(r"h([123])", name)
        if head:
            levels_seen.append((int(head.group(1)), paragraph.text.strip()[:20]))
    gate.checked += sum(len(items) for items in by_style.values())
    for name, items in by_style.items():
        sizes = {size for _, size, _ in items}
        fonts = {font for _, _, font in items}
        if len(sizes) > 1:
            gate.fail("style-size-mixed",
                      f"样式 {name} 在同一份文档里有多种字号：{'、'.join(str(s) for s in sorted(sizes))}pt",
                      items[0][0])
        if len(fonts) > 1:
            gate.fail("style-font-mixed", f"样式 {name} 混用了多种中文字体：{'、'.join(sorted(fonts))}",
                      items[0][0])
    previous = None
    for level, text in levels_seen:
        if previous is not None and level > previous + 1:
            gate.warn("heading-level-jump", f"标题层级从 h{previous} 跳到 h{level}：{text}")
        previous = level
    if direct:
        gate.warn("direct-formatting",
                  f"有 {direct} 处正文用了手工字号/字体（绕过样式）—— 改字号会漏改这些地方")


def gate_images(document, report: gates.GateReport) -> None:
    gate = report.gate("images", "图片（纵横比/尺寸/清晰度/替代文字）")
    text_width_cm = ((document.sections[0].page_width - document.sections[0].left_margin
                      - document.sections[0].right_margin) / 360000)
    from PIL import Image  # noqa: PLC0415

    for index, shape in enumerate(document.inline_shapes, start=1):
        gate.checked += 1
        width_cm = shape.width / 360000
        height_cm = shape.height / 360000
        if width_cm > text_width_cm + 0.05:
            gate.fail("image-too-wide",
                      f"图 {index} 宽 {width_cm:.1f}cm 超过版心 {text_width_cm:.1f}cm")
        try:
            embed = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            part = document.part.related_parts[embed]
            import io  # noqa: PLC0415

            with Image.open(io.BytesIO(part.blob)) as image:
                pixel_width, pixel_height = image.size
        except (KeyError, AttributeError, OSError):
            gate.warn("image-unreadable", f"图 {index} 读不出原始像素尺寸，跳过纵横比检查")
            continue
        intrinsic = pixel_width / pixel_height
        displayed = width_cm / height_cm
        if abs(displayed - intrinsic) / intrinsic > 0.05:
            gate.fail("image-distorted",
                      f"图 {index} 被拉伸：显示 {width_cm:.1f}×{height_cm:.1f}cm（比例 {displayed:.2f}）"
                      f"≠ 原始 {pixel_width}×{pixel_height}（{intrinsic:.2f}）")
        dpi = pixel_width / (width_cm / 2.54)
        if dpi < 96:
            gate.fail("image-low-dpi", f"图 {index} 有效分辨率 {dpi:.0f} dpi（印出来会糊）")
        elif dpi < 150:
            gate.warn("image-dpi", f"图 {index} 有效分辨率 {dpi:.0f} dpi（建议 ≥150）")
        descr = shape._inline.docPr.get("descr") if hasattr(shape._inline, "docPr") else ""
        if not (descr or "").strip():
            gate.warn("image-no-alt", f"图 {index} 没有替代文字（docx 无障碍与检索都靠它）")


def gate_toc(document, report: gates.GateReport) -> list[tuple[int, str, int | None]]:
    gate = report.gate("toc", "目录")
    entries = docx_fonts.read_toc_entries(document)
    body_xml = "".join(paragraph._p.xml for paragraph in document.paragraphs)
    has_field = "TOC \\o" in body_xml or 'TOC \\\\o' in body_xml
    gate.checked = max(len(entries), 1)
    if has_field and not entries:
        gate.notes.append("目录用的是 TOC 域：页码由 Word 更新（F9），静态无法校验；"
                          "建议用 --toc-mode static 生成可校验的静态目录")
        return entries
    if not entries:
        gate.fail("toc-missing", "文档没有目录（长篇文档应给目录）")
        return entries
    for level, text, page in entries:
        if page is None:
            gate.fail("toc-page-missing", f"目录条目没有页码：{text}")
    if len(entries) < 2:
        gate.warn("toc-thin", f"目录只有 {len(entries)} 条")
    return entries


def gate_page_setup(document, report: gates.GateReport) -> None:
    gate = report.gate("page-setup", "页面与页码")
    section = document.sections[0]
    width_cm = section.page_width / 360000
    height_cm = section.page_height / 360000
    gate.checked += 1
    if abs(width_cm - 21) > 0.3 or abs(height_cm - 29.7) > 0.3:
        gate.warn("paper-not-a4", f"纸张是 {width_cm:.1f}×{height_cm:.1f}cm（不是 A4）")
    for key, label in (("top", "上"), ("bottom", "下"), ("left", "左"), ("right", "右")):
        value = getattr(section, f"{key}_margin") / 360000
        gate.checked += 1
        if value < 1.5:
            gate.warn("margin-small", f"{label}页边距 {value:.2f}cm 偏小（正文会顶到纸边）")
        elif value > 4.0:
            gate.warn("margin-large", f"{label}页边距 {value:.2f}cm 偏大（版心太窄）")
    gate.checked += 1
    if not docx_fonts.footer_has_page_field(document):
        gate.fail("no-page-number", "页脚没有页码域（多页文档必须能看出第几页）")
    body_paragraphs = [p for p in document.paragraphs
                       if p.style is not None and p.style.name in ("body",)]
    if body_paragraphs:
        gate.checked += len(body_paragraphs)
        without_indent = [p for p in body_paragraphs
                          if not (p.paragraph_format.first_line_indent)]
        if len(without_indent) == len(body_paragraphs):
            gate.warn("no-first-line-indent", "正文段落都没有首行缩进（中文文档惯例是缩进 2 字符）")


def gate_hygiene(document, report: gates.GateReport, path: Path) -> None:
    gate = report.gate("hygiene", "交付卫生（路径/身份/凭据）")
    parts = hygiene.text_parts(path)
    findings: list[gates.Finding] = []
    for name, text in parts:
        findings += hygiene.scan_office_file_text(text, f"{path.name}:{name}")
    findings += hygiene.core_property_findings(document.core_properties)
    gate.checked = len(parts) + 1
    for finding in findings:
        gate.fail(finding.code, finding.message, finding.location)
    if not findings:
        gate.notes.append(f"扫过 {len(parts)} 个 XML 部件 + 核心属性：没有本机路径/用户名/凭据痕迹")


def gate_structure(document, report: gates.GateReport) -> None:
    gate = report.gate("structure", "结构卫生")
    gate.checked += 1
    if not [p for p in document.paragraphs if p.text.strip()] and not document.tables:
        gate.fail("empty-document", "文档没有任何内容")
        return
    trailing = 0
    for paragraph in reversed(document.paragraphs):
        if paragraph.text.strip():
            break
        trailing += 1
    gate.checked += 1
    if trailing > 2:
        gate.warn("trailing-empty", f"文末有 {trailing} 个空段落（导出后多出空白）")
    for index, table in enumerate(document.tables, start=1):
        gate.checked += 1
        if len(table.rows) > 25:
            header_xml = table.rows[0]._tr.xml
            if "tblHeader" not in header_xml:
                gate.fail("table-no-repeat-header",
                          f"表 {index} 有 {len(table.rows)} 行却未设跨页重复表头")
    toc_titles = {"目录", "目  录", "Contents"}
    in_toc = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text in toc_titles:
            in_toc = True
            continue
        if in_toc and "\t" in paragraph.text:
            continue
        if in_toc:
            in_toc = False
        if "\t" in paragraph.text and paragraph.text.strip():
            gate.checked += 1
            gate.warn("tab-alignment", f"段落用制表位排版（对齐会随字体变化跑偏）：{text[:30]}")


def gate_engine(path: Path, report: gates.GateReport, evidence_dir: Path | None,
                require_engine: bool, dpi: int, toc_entries: list[tuple[int, str, int | None]],
                toc_title: str) -> dict:
    engine = engines.engine_report()
    evidence: dict = {}
    gate = report.gate("render", "真渲染（逐页 PDF/PNG + 目录对账）")
    if not engine.can_convert:
        reason = ("本机没有 LibreOffice（soffice），真渲染与目录页码对账未执行。"
                  "装 LibreOffice 后跑 `python tools/env_doctor.py --office-home <安装目录>` 登记。")
        if require_engine:
            gate.add("fail", "engine-missing", reason)
        else:
            gate.skipped = reason
        return evidence
    workdir = evidence_dir or path.parent / f"{path.stem}.审查"
    workdir.mkdir(parents=True, exist_ok=True)
    try:
        pdf = engines.convert_to_pdf(path, workdir)
    except engines.ConvertError as error:
        gate.add("fail", "render-failed", str(error))
        return evidence
    evidence["pdf"] = str(pdf)
    pages = engines.pdf_text_per_page(pdf)
    gate.checked += len(pages)
    for index, text in enumerate(pages, start=1):
        if "□" in text or "\ufffd" in text:
            gate.fail("rendered-missing-glyph", "渲染结果里出现缺字形方框/替换符", f"第 {index} 页")
        for token in TOC_ERRORS:
            if token in text:
                gate.fail("toc-field-error", f"目录域报错：{token}", f"第 {index} 页")
        if len(text.strip()) < 20 and index not in (1,):
            gate.warn("nearly-blank-page", f"第 {index} 页几乎空白（分页问题？）")
    if toc_entries:
        problems = docx_toc.verify(pages, toc_entries, toc_title)
        gate.checked += len(toc_entries)
        for problem in problems:
            if problem["actual"] is None:
                gate.fail("toc-heading-unlocated",
                          f"目录里的标题在渲染结果里找不到：{problem['heading']}")
            else:
                gate.fail("toc-page-mismatch",
                          f"目录页码对不上：{problem['heading']} 写的是第 {problem['written']} 页，"
                          f"实际在第 {problem['actual']} 页")
        if not problems:
            gate.notes.append(f"目录 {len(toc_entries)} 条页码与渲染结果逐条一致")
    pngs = engines.pdf_to_png(pdf, workdir / "pages", dpi=dpi)
    evidence["pages"] = [str(item) for item in pngs]
    gate.notes.append(f"共 {len(pages)} 页，逐页图 {len(pngs)} 张 → {workdir / 'pages'}")
    return evidence


def main() -> int:
    cli.bootstrap()
    parser = argparse.ArgumentParser(description="docx 质量门")
    parser.add_argument("--docx", required=True)
    parser.add_argument("--evidence-dir", default="")
    parser.add_argument("--no-engine", action="store_true")
    parser.add_argument("--require-engine", action="store_true")
    parser.add_argument("--dpi", type=int, default=110)
    cli.add_common(parser)
    args = parser.parse_args()

    path = Path(args.docx).expanduser()
    if not path.is_file():
        print(f"找不到文件：{path}", file=sys.stderr)
        return 2
    report = gates.GateReport(subject=str(path))
    report.expected(EXPECTED_GATES)
    document = Document(str(path))
    toc_title = "目录"
    gate_font_size(document, report)
    gate_fonts(document, report)
    gate_style_consistency(document, report)
    gate_images(document, report)
    toc_entries = gate_toc(document, report)
    gate_page_setup(document, report)
    gate_hygiene(document, report, path)
    gate_structure(document, report)
    evidence: dict = {}
    if args.no_engine:
        report.gate("render", "真渲染（逐页 PDF/PNG + 目录对账）").skipped = "--no-engine：只跑静态判据"
    else:
        evidence = gate_engine(path, report,
                               Path(args.evidence_dir).expanduser() if args.evidence_dir else None,
                               args.require_engine, args.dpi, toc_entries, toc_title)
    report.meta["evidence"] = evidence
    payload = {"summary": report.render(), "report": report.to_dict()}
    cli.emit(payload, args)
    return report.exit_code


if __name__ == "__main__":
    raise SystemExit(main())
