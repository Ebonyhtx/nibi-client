"""读回 .docx 的结构与内容（JSON）：段落样式与字体、标题大纲、表格、图片、目录、页面设置。

    python tools/docx_read.py --docx 报告.docx [--text] [--json-out 结构.json]

与"打开看一眼"的区别：它把**样式与字体**（ascii/hAnsi/eastAsia 三套）、目录条目与页码、
页脚页码域、表格列数一致性都摊开 —— 交付缺陷大多藏在这些地方。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from nibi_kit import cli  # noqa: E402
from nibi_kit.env import require_modules  # noqa: E402

require_modules(["docx"], "docx_read（python-docx）")

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

import docx_fonts  # noqa: E402


def describe(document: Document) -> dict:
    paragraphs = []
    for index, paragraph in enumerate(document.paragraphs):
        if not paragraph.text.strip() and not paragraph.runs:
            continue
        fonts = docx_fonts.paragraph_fonts(paragraph, document)
        paragraphs.append({
            "index": index, "style": paragraph.style.name if paragraph.style else "",
            "text": paragraph.text[:200], "cjk": fonts["cjk"], "latin": fonts["latin"],
            "sizePt": fonts["size"], "align": str(paragraph.paragraph_format.alignment or ""),
        })
    tables = []
    for index, table in enumerate(document.tables):
        header = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
        tables.append({"index": index, "rows": len(table.rows),
                       "cols": len(table.columns), "header": header})
    images = []
    for index, shape in enumerate(document.inline_shapes):
        width_cm = round(shape.width / 360000, 2)
        height_cm = round(shape.height / 360000, 2)
        images.append({"index": index, "widthCm": width_cm, "heightCm": height_cm,
                       "ratio": round(width_cm / height_cm, 5) if height_cm else None})
    toc_entries = docx_fonts.read_toc_entries(document)
    section = document.sections[0]
    return {
        "paragraphCount": len(paragraphs),
        "paragraphs": paragraphs,
        "tables": tables,
        "images": images,
        "tocEntries": [{"level": level, "text": text, "page": page}
                       for level, text, page in toc_entries],
        "page": {"widthCm": round(section.page_width / 360000, 2),
                 "heightCm": round(section.page_height / 360000, 2),
                 "marginCm": {key: round(getattr(section, f"{key}_margin") / 360000, 2)
                              for key in ("top", "bottom", "left", "right")},
                 "footerHasPageField": docx_fonts.footer_has_page_field(document)},
        "coreProperties": {"title": document.core_properties.title or "",
                          "author": document.core_properties.author or "",
                          "lastModifiedBy": document.core_properties.last_modified_by or ""},
    }


def dump_text(document: Document) -> str:
    lines: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def main() -> int:
    cli.bootstrap()
    parser = argparse.ArgumentParser(description="读回 docx 结构与内容")
    parser.add_argument("--docx", required=True)
    parser.add_argument("--text", action="store_true", help="输出纯文本（段落 + 表格）")
    cli.add_common(parser)
    args = parser.parse_args()

    path = Path(args.docx).expanduser()
    if not path.is_file():
        print(f"找不到文件：{path}", file=sys.stderr)
        return 2
    document = Document(str(path))
    if args.text:
        text = dump_text(document)
        cli.emit({"summary": text, "text": text}, args)
        return 0
    payload = describe(document)
    lines = [f"{path.name}：{payload['paragraphCount']} 段、{len(payload['tables'])} 表、"
             f"{len(payload['images'])} 图、目录 {len(payload['tocEntries'])} 条",
             f"  页面：{payload['page']['widthCm']}×{payload['page']['heightCm']}cm，"
             f"页边距 上{payload['page']['marginCm']['top']}/下{payload['page']['marginCm']['bottom']}/"
             f"左{payload['page']['marginCm']['left']}/右{payload['page']['marginCm']['right']}cm，"
             f"页脚页码域：{'有' if payload['page']['footerHasPageField'] else '没有'}"]
    for entry in payload["tocEntries"]:
        lines.append(f"  · 目录[{entry['level']}] {entry['text']} → 第 {entry['page']} 页")
    for item in payload["images"]:
        lines.append(f"  · 图 {item['index'] + 1}：{item['widthCm']}×{item['heightCm']}cm")
    cli.emit({"summary": "\n".join(lines), **payload}, args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
