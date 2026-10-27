"""质量门冒烟：正向对照 + 反向对照，防止判据"静默空转"。

    python tools/docx_smoke.py [--keep]

1. **正向对照**：用 docx_build 造一份正常文档 → 质量门必须通过（比了 N 项，N>0）。
   需要 LibreOffice 才能定目录页码；没有引擎时这一项**显式跳过并说明原因**。
2. **反向对照**：用 python-docx 直接造一份**逐条犯规**的文档（6pt 字号、中文配拉丁字体、
   图片被拉扁、标题跳级、没有页码、没有目录、正文里带本机路径）→ 质量门必须失败，
   且每个期望判据码都要出现。缺任何一条 → 退出码非零。

反向对照是重点：门禁全绿有两种可能 —— 东西真好，或者判据压根没跑。
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from nibi_kit import cli, engines  # noqa: E402
from nibi_kit.env import require_modules  # noqa: E402

require_modules(["docx", "PIL"], "docx_smoke（python-docx + pillow）")

PLUGIN_ROOT = Path(__file__).resolve().parent.parent
# 自测夹具**不随插件发布**（用户不该在安装目录里看到"样例报告"）：
# 它们只在仓库里（templates/_dev/），用于开发与打包前验收。
DEV_FIXTURES = PLUGIN_ROOT / "templates" / "_dev"
TOOLS = Path(__file__).resolve().parent

EXPECTED_CODES = {
    "font-too-small",       # 6pt 字号
    "font-not-cjk",         # 中文配 Calibri（Word 会替换字体）
    "image-distorted",      # 图片被拉扁
    "heading-level-jump",   # h1 → h3 跳级
    "no-page-number",       # 页脚没有页码域
    "toc-missing",          # 有标题却没有目录
    "win-drive-path",       # 正文里带本机路径
}


def run_tool(script: str, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(TOOLS / script), *args],  # noqa: S603
                          capture_output=True, text=True, encoding="utf-8", errors="replace")


def load_report(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    return data.get("report", data)


def positive_control(workdir: Path) -> tuple[str, str]:
    spec = DEV_FIXTURES / "spec-report.json"
    if not spec.is_file():
        return "skip", ("正向对照需要仓库内的样板（templates/_dev/spec-report.json）；"
                        "该夹具不随插件发布，安装目录里跑本项会跳过。")
    target = workdir / "正向对照.docx"
    built = run_tool("docx_build.py", "--spec", str(spec), "--out", str(target))
    if built.returncode != 0 or not target.is_file():
        return "fail", f"构建失败：{built.stderr.strip()[-300:]}"
    report_path = workdir / "正向对照.判据.json"
    gated = run_tool("docx_gate.py", "--docx", str(target), "--no-engine",
                     "--json-out", str(report_path))
    if not report_path.is_file():
        return "fail", "质量门没有产出报告"
    report = load_report(report_path)
    ran = [gate for gate in report["gates"] if gate["checked"] > 0]
    if gated.returncode != 0:
        bad = [item for gate in report["gates"] for item in gate["findings"]
               if item["severity"] == "fail"]
        return "fail", "正常文档被拦下：" + json.dumps(bad[:4], ensure_ascii=False)
    if len(ran) < 7:
        return "fail", f"只有 {len(ran)} 条判据真的比过东西（应有 ≥7）"
    return "ok", f"{len(ran)} 条判据共比了 {sum(gate['checked'] for gate in ran)} 项"


def make_broken(workdir: Path) -> Path:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    from PIL import Image

    import docx_build
    import docx_spec

    target = workdir / "反向对照.docx"
    document = Document()
    docx_build.build_styles(document, docx_spec.styles_with_overrides(None),
                            docx_spec.palette("nibi"))
    document.add_paragraph("一、项目概述", style="h1")
    small = document.add_paragraph("这一段用了六号字，印出来看不清。", style="body")
    small.runs[0].font.size = Pt(6)
    latin = document.add_paragraph("中文配拉丁字体：", style="body")
    run = latin.add_run("这段中文的 eastAsia 被写成 Calibri，Word 会替换字体。")
    rpr = run._element.get_or_add_rPr()
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "Calibri")
    rpr.append(fonts)
    document.add_paragraph("三、结论", style="h3")      # 从 h1 跳到 h3
    image_path = workdir / "broken.png"
    Image.new("RGB", (800, 600), "white").save(image_path)
    document.add_picture(str(image_path), width=Cm(6), height=Cm(3))   # 4:3 的图被拉成 2:1
    document.add_paragraph("C:\\Users\\<用户名>\\报告.docx", style="body")
    document.save(target)
    return target


def negative_control(workdir: Path) -> tuple[str, str]:
    target = make_broken(workdir)
    report_path = workdir / "反向对照.判据.json"
    gated = run_tool("docx_gate.py", "--docx", str(target), "--no-engine",
                     "--json-out", str(report_path))
    if not report_path.is_file():
        return "fail", "质量门没有产出报告"
    report = load_report(report_path)
    codes = {item["code"] for gate in report["gates"] for item in gate["findings"]}
    if gated.returncode == 0:
        return "fail", "犯规文档竟然通过了质量门（判据空转）"
    missing = sorted(EXPECTED_CODES - codes)
    if missing:
        return "fail", "以下判据没有抓到对应犯规：\n" + "\n".join(f"{code}（期望命中）"
                                                                 for code in missing)
    return "ok", f"{len(EXPECTED_CODES)} 条期望判据全部命中，另有 {len(codes - EXPECTED_CODES)} 条附带命中"


def main() -> int:
    cli.bootstrap()
    parser = argparse.ArgumentParser(description="质量门冒烟（正向 + 反向对照）")
    parser.add_argument("--keep", action="store_true")
    cli.add_common(parser)
    args = parser.parse_args()

    engine = engines.engine_report()
    with tempfile.TemporaryDirectory(prefix="nibi-docx-smoke-") as raw:
        workdir = Path(raw)
        results: list[tuple[str, str, str]] = []
        if engine.can_convert:
            status, detail = positive_control(workdir)
        else:
            status, detail = "skip", "没有 LibreOffice：目录页码无法解析，正向对照本轮跳过（不是通过）"
        results.append(("正向对照（正常文档必须过）", status, detail))
        status, detail = negative_control(workdir)
        results.append(("反向对照（犯规文档必须被抓）", status, detail))

        lines = ["docx 质量门冒烟："]
        for title, state, detail in results:
            mark = {"ok": "ok  ", "fail": "FAIL", "skip": "skip"}[state]
            lines.append(f"  [{mark}] {title}")
            for chunk in detail.splitlines():
                lines.append(f"         {chunk}")
        failed = [title for title, state, _ in results if state == "fail"]
        skipped = [title for title, state, _ in results if state == "skip"]
        conclusion = "通过" if not failed and not skipped else (
            "未通过——" + "、".join(failed) if failed else f"部分未验（{len(skipped)} 项跳过）")
        lines.append(f"  结论：{conclusion}")
        if args.keep and (failed or skipped):
            keep = Path.cwd() / "smoke-failure"
            keep.mkdir(exist_ok=True)
            for item in workdir.iterdir():
                if item.is_file():
                    (keep / item.name).write_bytes(item.read_bytes())
            lines.append(f"  （现场已保留到 {keep}）")
        cli.emit({"summary": "\n".join(lines),
                  "results": [{"case": title, "status": state, "detail": detail}
                              for title, state, detail in results]}, args)
        return 0 if not failed else 1


if __name__ == "__main__":
    raise SystemExit(main())
