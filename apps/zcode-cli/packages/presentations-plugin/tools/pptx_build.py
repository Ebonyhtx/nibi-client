"""把 JSON 规格落成 .pptx（python-pptx）：16:9、空白版式 + 显式文本框、投影字号基线。

用法：
    python tools/pptx_build.py --spec 演示.json --out 交付/答辩.pptx

设计要点：

* **空白版式 + 显式文本框**：不用占位符，避免"实际字号取决于母版"这类判据测不到的情况；
* **中文字体要写 `a:ea`**（DrawingML 里 `a:latin` 只管拉丁字符，中文走 `a:ea`）——
  只设 latin 的话中文会用主题字体，同一份 PPT 在不同机器上长得不一样；
* **图片只给一个维度**（按原始纵横比缩放到内容框内），杜绝被拉扁；
* **每页都过一遍版心检查**：文本框按内容行数给高度，不靠"看起来差不多"。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from nibi_kit import cli, hygiene  # noqa: E402
from nibi_kit import textmetrics  # noqa: E402
from nibi_kit.env import require_modules  # noqa: E402

require_modules(["pptx", "PIL"], "pptx_build（python-pptx + pillow）")

import pptx_spec as spec_mod  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
CHART_TYPES = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "barStacked": XL_CHART_TYPE.COLUMN_STACKED,
               "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE,
               "doughnut": XL_CHART_TYPE.DOUGHNUT}
MARGIN_IN = 0.62
TITLE_TOP_IN = 0.35
TITLE_HEIGHT_IN = 0.95
CONTENT_TOP_IN = 1.45
CONTENT_BOTTOM_IN = 6.85
FOOTER_TOP_IN = 6.95


class BuildError(RuntimeError):
    pass


# ---------------------------------------------------------------- 文字


def style_of(styles: dict, name: str) -> dict:
    if name not in styles:
        raise spec_mod.SpecError(f"未知样式名：{name}（可用：{'、'.join(sorted(styles))}）")
    return styles[name]


def set_run_font(run, definition: dict, palette: dict) -> None:
    """字号/粗细/颜色 + **中文字体（a:ea）**。只设 a:latin 时中文走主题字体。"""
    size = float(definition.get("size", 18))
    run.font.size = Pt(size)
    run.font.bold = bool(definition.get("bold", False))
    run.font.italic = bool(definition.get("italic", False))
    color = definition.get("color")
    if color:
        run.font.color.rgb = RGBColor.from_string(palette.get(color, str(color).lstrip("#")))
    family = str(definition.get("cjk", "微软雅黑"))
    run.font.name = family
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", family)


def add_textbox(slide, x: float, y: float, width: float, height: float, *, anchor=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    if anchor is not None:
        frame.vertical_anchor = anchor
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.0)
    frame.margin_bottom = Inches(0.0)
    return shape


def fill_paragraph(paragraph, text: str, definition: dict, palette: dict, *,
                   align: str = "left", space_after: float = 6.0, bullet: str = "") -> None:
    paragraph.alignment = ALIGN.get(align, PP_ALIGN.LEFT)
    if definition.get("lineSpacing"):
        paragraph.line_spacing = float(definition["lineSpacing"])
    paragraph.space_after = Pt(space_after)
    run = paragraph.add_run()
    run.text = text
    set_run_font(run, definition, palette)
    if bullet:
        apply_bullet(paragraph, bullet, float(definition.get("size", 18)))


def apply_bullet(paragraph, char: str, size_pt: float) -> None:
    """真项目符号（`a:buChar`）+ 悬挂缩进；不要靠手写 "• " 前缀（换行会缩进错位）。"""
    ppr = paragraph._p.get_or_add_pPr()
    # 悬挂缩进：左缩进 0.32in（292608 EMU），首行回退同样的距离
    ppr.set("marL", "292608")
    ppr.set("indent", "-292608")
    for tag, attrs in (("a:buFont", {"typeface": "Arial"}), ("a:buChar", {"char": char})):
        node = ppr.makeelement(qn(tag), attrs)
        ppr.append(node)


def text_height_in(lines: int, size_pt: float, line_spacing: float = 1.25) -> float:
    return lines * size_pt * line_spacing / 72.0 + 0.06


def content_width_in() -> float:
    return spec_mod.SLIDE_WIDTH_IN - 2 * MARGIN_IN


def estimate_lines(text: str, size_pt: float, width_in: float, family: str = "微软雅黑") -> int:
    return textmetrics.estimate_lines(text, family, size_pt,
                                      textmetrics.column_width_px(width_in * 96 / 7.0) - 6)


# ---------------------------------------------------------------- 通用零件


def slide_title(slide, text: str, styles: dict, palette: dict,
                width: float | None = None) -> float:
    width = width or content_width_in()
    definition = style_of(styles, "slideTitle")
    lines = estimate_lines(text, float(definition["size"]), width)
    height = max(TITLE_HEIGHT_IN, text_height_in(lines, float(definition["size"]), 1.1))
    shape = add_textbox(slide, MARGIN_IN, TITLE_TOP_IN, width, height)
    fill_paragraph(shape.text_frame.paragraphs[0], text, definition, palette)
    underline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN_IN),
                                       Inches(TITLE_TOP_IN + height + 0.05),
                                       Inches(1.5), Pt(3))
    underline.fill.solid()
    underline.fill.fore_color.rgb = RGBColor.from_string(palette["orange"])
    underline.line.fill.background()
    underline.shadow.inherit = False
    return TITLE_TOP_IN + height + 0.22


def add_page_number(slide, index: int, styles: dict, palette: dict) -> None:
    definition = style_of(styles, "pageNumber")
    shape = add_textbox(slide, spec_mod.SLIDE_WIDTH_IN - MARGIN_IN - 1.0, FOOTER_TOP_IN, 1.0, 0.3)
    # 命名让判据能认出"版式件"（页码不受正文字号软下限约束）
    shape.name = "pageNumber"
    fill_paragraph(shape.text_frame.paragraphs[0], str(index), definition, palette, align="right")


def add_notes(slide, text: str) -> None:
    if not text:
        return
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- 各版式


def build_title_slide(deck, slide_spec: dict, styles: dict, palette: dict) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                  Inches(spec_mod.SLIDE_WIDTH_IN), Inches(2.55))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string(palette["navy"])
    band.line.fill.background()
    band.shadow.inherit = False
    definition = style_of(styles, "sectionTitle")
    shape = add_textbox(slide, MARGIN_IN, 0.95, content_width_in(), 1.2, anchor=MSO_ANCHOR.MIDDLE)
    fill_paragraph(shape.text_frame.paragraphs[0], str(slide_spec.get("title", "")),
                   definition, palette)
    subtitle = style_of(styles, "sectionSubtitle")
    shape = add_textbox(slide, MARGIN_IN, 2.85, content_width_in(), 1.6)
    lines = slide_spec.get("subtitle", [])
    lines = [lines] if isinstance(lines, str) else list(lines)
    for index, line in enumerate(lines):
        paragraph = shape.text_frame.paragraphs[0] if index == 0 else shape.text_frame.add_paragraph()
        fill_paragraph(paragraph, str(line), subtitle, palette, space_after=4)


def build_section_slide(deck, slide_spec: dict, styles: dict, palette: dict) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                        Inches(spec_mod.SLIDE_WIDTH_IN),
                                        Inches(spec_mod.SLIDE_HEIGHT_IN))
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor.from_string(palette["navyDeep"])
    background.line.fill.background()
    background.shadow.inherit = False
    definition = style_of(styles, "sectionTitle")
    shape = add_textbox(slide, MARGIN_IN, 2.6, content_width_in(), 1.3, anchor=MSO_ANCHOR.MIDDLE)
    fill_paragraph(shape.text_frame.paragraphs[0], str(slide_spec.get("title", "")),
                   definition, palette)
    if slide_spec.get("subtitle"):
        subtitle = style_of(styles, "sectionSubtitle")
        shape = add_textbox(slide, MARGIN_IN, 3.95, content_width_in(), 0.6)
        fill_paragraph(shape.text_frame.paragraphs[0], str(slide_spec["subtitle"]), subtitle, palette)


def build_bullets_slide(deck, slide_spec: dict, styles: dict, palette: dict, index: int) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    top = slide_title(slide, str(slide_spec.get("title", "")), styles, palette)
    definition = style_of(styles, slide_spec.get("style", "bulletText"))
    items = slide_spec.get("items", [])
    height = CONTENT_BOTTOM_IN - top
    shape = add_textbox(slide, MARGIN_IN, top, content_width_in(), height)
    frame = shape.text_frame
    for item_index, item in enumerate(items):
        text = spec_mod.text_of(item)
        level = int(item.get("level", 0)) if isinstance(item, dict) else 0
        paragraph = frame.paragraphs[0] if item_index == 0 else frame.add_paragraph()
        fill_paragraph(paragraph, text, definition, palette,
                       bullet="•" if level == 0 else "–", space_after=10 - 2 * level)
        if level:
            paragraph._p.get_or_add_pPr().set("marL", str(int(0.32 * 914400 * (1 + level))))
    add_page_number(slide, index, styles, palette)


def build_two_column_slide(deck, slide_spec: dict, styles: dict, palette: dict, index: int) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    top = slide_title(slide, str(slide_spec.get("title", "")), styles, palette)
    gap = 0.35
    column_width = (content_width_in() - gap) / 2
    heading_style = style_of(styles, "cardTitle")
    body_style = style_of(styles, slide_spec.get("style", "bodyTextSmall"))
    for column_index, column in enumerate(slide_spec.get("columns", [])[:2]):
        x = MARGIN_IN + column_index * (column_width + gap)
        heading = str(column.get("heading", ""))
        shape = add_textbox(slide, x, top, column_width, 0.45)
        fill_paragraph(shape.text_frame.paragraphs[0], heading, heading_style, palette, space_after=2)
        body = add_textbox(slide, x, top + 0.5, column_width, CONTENT_BOTTOM_IN - top - 0.5)
        frame = body.text_frame
        lines = column.get("lines", [])
        for line_index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
            fill_paragraph(paragraph, spec_mod.text_of(line), body_style, palette,
                           bullet="•", space_after=6)
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(top + 0.42),
                                      Inches(column_width), Pt(1.5))
        rule.fill.solid()
        rule.fill.fore_color.rgb = RGBColor.from_string(palette["line"])
        rule.line.fill.background()
        rule.shadow.inherit = False
    add_page_number(slide, index, styles, palette)


def build_cards_slide(deck, slide_spec: dict, styles: dict, palette: dict, index: int) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    top = slide_title(slide, str(slide_spec.get("title", "")), styles, palette)
    cards = slide_spec.get("cards", [])[:4]
    gap = 0.3
    width = (content_width_in() - gap * (len(cards) - 1)) / max(len(cards), 1)
    # 高度取"剩余空间的一半以上并居中"：卡片贴顶会让下半页空掉（视觉失衡）
    available = CONTENT_BOTTOM_IN - top - 0.2
    height = min(2.8, max(1.9, available * 0.62))
    card_top = top + max(0.0, (available - height) / 2)
    for card_index, card in enumerate(cards):
        x = MARGIN_IN + card_index * (width + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(card_top),
                                     Inches(width), Inches(height))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string(palette.get(card.get("fill", "cream"), "F7EFE2"))
        box.line.color.rgb = RGBColor.from_string(palette["line"])
        box.shadow.inherit = False
        box.text_frame.text = ""
        value_style = style_of(styles, "cardValue")
        label_style = style_of(styles, "cardLabel")
        frame = box.text_frame
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        fill_paragraph(paragraph, spec_mod.text_of(card.get("value", "")), value_style, palette,
                       align="center", space_after=2)
        for extra in [card.get("label", "")] + list(card.get("lines", [])):
            if not extra:
                continue
            paragraph = frame.add_paragraph()
            fill_paragraph(paragraph, spec_mod.text_of(extra), label_style, palette,
                           align="center", space_after=2)
    if slide_spec.get("note"):
        note = style_of(styles, "sourceNote")
        shape = add_textbox(slide, MARGIN_IN, CONTENT_BOTTOM_IN - 0.45, content_width_in(), 0.4)
        shape.name = "sourceNote"
        fill_paragraph(shape.text_frame.paragraphs[0], str(slide_spec["note"]), note, palette)
    add_page_number(slide, index, styles, palette)


def build_table_slide(deck, slide_spec: dict, styles: dict, palette: dict, index: int) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    top = slide_title(slide, str(slide_spec.get("title", "")), styles, palette)
    header = [str(item) for item in slide_spec.get("header", [])]
    rows = slide_spec.get("rows", [])
    if not header or not rows:
        raise BuildError("table 页需要 header 与 rows")
    height = min(CONTENT_BOTTOM_IN - top, 0.5 + 0.42 * (len(rows) + 1))
    table_shape = slide.shapes.add_table(len(rows) + 1, len(header), Inches(MARGIN_IN),
                                         Inches(top), Inches(content_width_in()), Inches(height))
    table = table_shape.table
    header_style = style_of(styles, "tableHeader")
    text_style = style_of(styles, "tableText")
    for column_index, text in enumerate(header):
        cell = table.cell(0, column_index)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(palette["navy"])
        paragraph = cell.text_frame.paragraphs[0]
        fill_paragraph(paragraph, text, header_style, palette, align="center", space_after=0)
    numeric = [all(str(row[column_index]).replace(",", "").replace("%", "").replace("+", "")
                   .replace(".", "").replace("-", "").strip().isdigit()
                   for row in rows if column_index < len(row)) for column_index in range(len(header))]
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row[:len(header)]):
            cell = table.cell(row_index, column_index)
            paragraph = cell.text_frame.paragraphs[0]
            fill_paragraph(paragraph, str(value), text_style, palette,
                           align="right" if numeric[column_index] and column_index > 0 else "left",
                           space_after=0)
    if slide_spec.get("caption"):
        caption = style_of(styles, "sourceNote")
        shape = add_textbox(slide, MARGIN_IN, top + height + 0.1, content_width_in(), 0.4)
        shape.name = "sourceNote"
        fill_paragraph(shape.text_frame.paragraphs[0], str(slide_spec["caption"]), caption, palette)
    add_page_number(slide, index, styles, palette)


def build_chart_slide(deck, slide_spec: dict, styles: dict, palette: dict, index: int) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    top = slide_title(slide, str(slide_spec.get("title", "")), styles, palette)
    chart_spec = slide_spec.get("chart", {})
    data = CategoryChartData()
    data.categories = [str(item) for item in chart_spec.get("categories", [])]
    for series in chart_spec.get("series", []):
        data.add_series(str(series.get("name", "")), tuple(series.get("values", [])))
    kind = CHART_TYPES.get(str(chart_spec.get("type", "bar")))
    if kind is None:
        raise spec_mod.SpecError(f"未知图表类型：{chart_spec.get('type')}（可选：{'、'.join(CHART_TYPES)}）")
    height = min(CONTENT_BOTTOM_IN - top - 0.5, 4.6)
    frame = slide.shapes.add_chart(kind, Inches(MARGIN_IN), Inches(top),
                                   Inches(content_width_in()), Inches(height), data)
    chart = frame.chart
    chart.has_title = bool(chart_spec.get("title"))
    if chart_spec.get("title"):
        chart.chart_title.text_frame.text = str(chart_spec["title"])
    chart.font.size = Pt(12)
    chart.font.name = "微软雅黑"
    series_colors = [palette["navy"], palette["orange"], palette["navySoft"],
                     palette["positive"], palette["negative"]]
    for series_index, series in enumerate(chart.series):
        color = series_colors[series_index % len(series_colors)]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
        if kind in (XL_CHART_TYPE.LINE_MARKERS,):
            series.format.line.color.rgb = RGBColor.from_string(color)
    if kind in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT):
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
    else:
        chart.has_legend = bool(chart_spec.get("legend", True))
        if chart.legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
    if slide_spec.get("note"):
        note = style_of(styles, "sourceNote")
        shape = add_textbox(slide, MARGIN_IN, CONTENT_BOTTOM_IN - 0.45, content_width_in(), 0.4)
        shape.name = "sourceNote"
        fill_paragraph(shape.text_frame.paragraphs[0], str(slide_spec["note"]), note, palette)
    add_page_number(slide, index, styles, palette)


def build_image_slide(deck, slide_spec: dict, styles: dict, palette: dict, index: int,
                      spec_dir: Path | None) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    top = slide_title(slide, str(slide_spec.get("title", "")), styles, palette)
    raw = str(slide_spec.get("path", ""))
    path = Path(raw).expanduser()
    if not path.is_absolute() and spec_dir is not None:
        path = (spec_dir / raw).resolve()
    if not path.is_file():
        raise BuildError(f"image 找不到图片：{path}")
    from PIL import Image  # noqa: PLC0415

    with Image.open(path) as image:
        pixel_width, pixel_height = image.size
    box_width = content_width_in()
    box_height = CONTENT_BOTTOM_IN - top - (0.5 if slide_spec.get("caption") else 0.1)
    ratio = pixel_width / pixel_height
    width = box_width
    height = width / ratio
    if height > box_height:      # 只缩一个维度，保持原始纵横比
        height = box_height
        width = height * ratio
    left = MARGIN_IN + (box_width - width) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))
    if slide_spec.get("caption"):
        caption = style_of(styles, "caption")
        shape = add_textbox(slide, MARGIN_IN, top + height + 0.12, content_width_in(), 0.4)
        shape.name = "sourceNote"
        fill_paragraph(shape.text_frame.paragraphs[0], str(slide_spec["caption"]), caption,
                       palette, align="center")
        alt = str(slide_spec.get("alt", "")).strip()
        if alt:
            try:
                slide.shapes[-1]._element._nvXxPr.cNvPr.set("descr", alt)
            except AttributeError:
                pass
    add_page_number(slide, index, styles, palette)


def build_quote_slide(deck, slide_spec: dict, styles: dict, palette: dict, index: int) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    definition = dict(style_of(styles, "sectionTitle"))
    definition["size"] = float(slide_spec.get("size", 26))
    definition["color"] = slide_spec.get("color", "navy")
    shape = add_textbox(slide, MARGIN_IN + 0.8, 2.4, content_width_in() - 1.6, 2.0,
                        anchor=MSO_ANCHOR.MIDDLE)
    fill_paragraph(shape.text_frame.paragraphs[0], str(slide_spec.get("text", "")),
                   definition, palette, align="center")
    if slide_spec.get("source"):
        source = style_of(styles, "slideSubtitle")
        shape = add_textbox(slide, MARGIN_IN, 4.6, content_width_in(), 0.5)
        fill_paragraph(shape.text_frame.paragraphs[0], str(slide_spec["source"]), source,
                       palette, align="right")
    add_page_number(slide, index, styles, palette)


def build_end_slide(deck, slide_spec: dict, styles: dict, palette: dict) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                  Inches(spec_mod.SLIDE_WIDTH_IN), Inches(spec_mod.SLIDE_HEIGHT_IN))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string(palette["navy"])
    band.line.fill.background()
    band.shadow.inherit = False
    definition = style_of(styles, "sectionTitle")
    shape = add_textbox(slide, MARGIN_IN, 3.0, content_width_in(), 1.0, anchor=MSO_ANCHOR.MIDDLE)
    fill_paragraph(shape.text_frame.paragraphs[0], str(slide_spec.get("title", "谢谢")),
                   definition, palette, align="center")
    if slide_spec.get("subtitle"):
        subtitle = style_of(styles, "sectionSubtitle")
        shape = add_textbox(slide, MARGIN_IN, 4.1, content_width_in(), 0.6)
        fill_paragraph(shape.text_frame.paragraphs[0], str(slide_spec["subtitle"]), subtitle,
                       palette, align="center")


# ---------------------------------------------------------------- 主流程


def build(spec: dict, out_path: Path, spec_dir: Path | None = None) -> dict:
    deck_spec = spec.get("deck", {})
    styles = spec_mod.styles_with_overrides(spec.get("styles"))
    palette = spec_mod.palette(deck_spec.get("theme", "nibi"))

    deck = Presentation()
    deck.slide_width = Inches(spec_mod.SLIDE_WIDTH_IN)
    deck.slide_height = Inches(spec_mod.SLIDE_HEIGHT_IN)
    deck.core_properties.title = str(deck_spec.get("title", ""))
    deck.core_properties.author = str(deck_spec.get("author", "NIBI"))
    hygiene.scrub_core_properties(deck.core_properties, author=str(deck_spec.get("author", "NIBI")))

    summary = {"slides": 0, "kinds": {}, "images": 0, "charts": 0, "tables": 0, "textboxes": 0}
    page_index = 0
    page_created = 0
    for slide_spec in spec["slides"]:
        kind = str(slide_spec.get("kind", "bullets"))
        page_index += 1
        if kind == "title":
            build_title_slide(deck, slide_spec, styles, palette)
        elif kind == "section":
            build_section_slide(deck, slide_spec, styles, palette)
            page_index -= 1
        elif kind == "end":
            build_end_slide(deck, slide_spec, styles, palette)
            page_index -= 1
        elif kind == "bullets":
            build_bullets_slide(deck, slide_spec, styles, palette, page_index)
        elif kind == "twoColumn":
            build_two_column_slide(deck, slide_spec, styles, palette, page_index)
        elif kind == "cards":
            build_cards_slide(deck, slide_spec, styles, palette, page_index)
        elif kind == "table":
            build_table_slide(deck, slide_spec, styles, palette, page_index)
            summary["tables"] += 1
        elif kind == "chart":
            build_chart_slide(deck, slide_spec, styles, palette, page_index)
            summary["charts"] += 1
        elif kind == "image":
            build_image_slide(deck, slide_spec, styles, palette, page_index, spec_dir)
            summary["images"] += 1
        elif kind == "quote":
            build_quote_slide(deck, slide_spec, styles, palette, page_index)
        else:
            raise BuildError(f"未知版式：{kind}")
        summary["kinds"][kind] = summary["kinds"].get(kind, 0) + 1
        page_created += 1
        slide = deck.slides[-1]
        summary["textboxes"] += sum(1 for shape in slide.shapes if shape.has_text_frame)
        add_notes(slide, str(slide_spec.get("notes", "")))
    summary["slides"] = page_created

    out_path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(out_path)
    return summary


def main() -> int:
    cli.bootstrap()
    parser = argparse.ArgumentParser(description="从 JSON 规格构建 .pptx")
    parser.add_argument("--spec", required=True)
    parser.add_argument("--out", required=True)
    cli.add_common(parser)
    args = parser.parse_args()

    try:
        spec = spec_mod.load_spec(args.spec)
    except spec_mod.SpecError as error:
        print(f"规格不合法：{error}", file=sys.stderr)
        return 2
    out_path = Path(args.out).expanduser()
    try:
        summary = build(spec, out_path, Path(args.spec).expanduser().resolve().parent)
    except (BuildError, spec_mod.SpecError) as error:
        print(f"构建失败：{error}", file=sys.stderr)
        return 2
    size_kb = out_path.stat().st_size / 1024 if out_path.is_file() else 0
    kinds = "、".join(f"{name}×{count}" for name, count in sorted(summary["kinds"].items()))
    lines = [f"已生成 {out_path}（{size_kb:.1f} KB）",
             f"  共 {summary['slides']} 页（{kinds}）",
             f"  表格 {summary['tables']}、图表 {summary['charts']}、图片 {summary['images']}、"
             f"文本框 {summary['textboxes']}",
             "下一步：python tools/pptx_gate.py --pptx <文件>；python tools/pptx_render.py 出逐页图"]
    cli.emit({"summary": "\n".join(lines), "out": str(out_path), "build": summary}, args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
