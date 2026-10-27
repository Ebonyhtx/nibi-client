"""读回 .pptx 的结构与文字（JSON）：页面、形状、文本框、字号、表格、图片、备注。

    python tools/pptx_read.py --pptx 答辩.pptx [--text] [--json-out 结构.json]

关注点与"打开翻一遍"不同：它把**每个文本框的几何与字号**摊开（越界与溢出判据就靠这些），
外加图片的显示尺寸与原始像素（纵横比是否被拉）与备注（答辩备注经常被忘在规格里没进文件）。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from nibi_kit import cli  # noqa: E402
from nibi_kit.env import require_modules  # noqa: E402

require_modules(["pptx"], "pptx_read（python-pptx）")

from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

EMU_PER_IN = 914400


def shape_info(shape) -> dict:
    info = {"name": shape.shape_type and str(shape.shape_type), "leftIn": round(shape.left / EMU_PER_IN, 2),
            "topIn": round(shape.top / EMU_PER_IN, 2),
            "widthIn": round(shape.width / EMU_PER_IN, 2),
            "heightIn": round(shape.height / EMU_PER_IN, 2)}
    if shape.has_text_frame:
        sizes = []
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text.strip():
                    sizes.append(run.font.size.pt if run.font.size else None)
        info["text"] = shape.text_frame.text[:160]
        info["fontSizes"] = sizes
        info["wordWrap"] = bool(shape.text_frame.word_wrap)
    if shape.has_table:
        table = shape.table
        info["table"] = {"rows": len(table.rows), "cols": len(table.columns),
                         "header": [cell.text for cell in table.rows[0].cells]}
    if shape.shape_type == 13 or getattr(shape, "image", None) is not None:
        try:
            info["image"] = {"widthIn": info["widthIn"], "heightIn": info["heightIn"],
                             "pixelWidth": shape.image.size[0], "pixelHeight": shape.image.size[1]}
        except AttributeError:
            pass
    if getattr(shape, "has_chart", False):
        info["chart"] = {"type": str(shape.chart.chart_type)}
    return info


def describe(deck) -> dict:
    slides = []
    for index, slide in enumerate(deck.slides, start=1):
        shapes = [shape_info(shape) for shape in slide.shapes]
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        title = ""
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                title = shape.text_frame.text.strip().splitlines()[0]
                break
        slides.append({"index": index, "title": title[:60], "shapes": shapes,
                       "notes": notes[:200], "shapeCount": len(shapes)})
    return {"slideWidthIn": round(deck.slide_width / EMU_PER_IN, 3),
            "slideHeightIn": round(deck.slide_height / EMU_PER_IN, 3),
            "slideCount": len(slides), "slides": slides}


def dump_text(deck) -> str:
    lines: list[str] = []
    for index, slide in enumerate(deck.slides, start=1):
        lines.append(f"--- 第 {index} 页 ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append(f"[备注] {slide.notes_slide.notes_text_frame.text.strip()}")
    return "\n".join(lines)


def main() -> int:
    cli.bootstrap()
    parser = argparse.ArgumentParser(description="读回 pptx 结构")
    parser.add_argument("--pptx", required=True)
    parser.add_argument("--text", action="store_true")
    cli.add_common(parser)
    args = parser.parse_args()

    path = Path(args.pptx).expanduser()
    if not path.is_file():
        print(f"找不到文件：{path}", file=sys.stderr)
        return 2
    deck = Presentation(str(path))
    if args.text:
        text = dump_text(deck)
        cli.emit({"summary": text, "text": text}, args)
        return 0
    payload = describe(deck)
    sizes = sorted({size for slide in payload["slides"] for shape in slide["shapes"]
                    for size in shape.get("fontSizes", []) if size})
    lines = [f"{path.name}：{payload['slideCount']} 页，"
             f"{payload['slideWidthIn']}×{payload['slideHeightIn']}in",
             f"  用到的字号：{'、'.join(f'{size:g}pt' for size in sizes) or '（无）'}"]
    for slide in payload["slides"]:
        notes = "，含备注" if slide["notes"] else ""
        lines.append(f"  · 第 {slide['index']} 页：{slide['title']}（{slide['shapeCount']} 个形状{notes}）")
    cli.emit({"summary": "\n".join(lines), **payload}, args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
