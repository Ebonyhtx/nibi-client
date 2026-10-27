"""出逐页图给视觉验收用（visual-judge 或人眼）。

    python tools/pptx_render.py --pptx 答辩.pptx --outdir 审查/演示图 [--dpi 110]

* 有 LibreOffice → 真渲染：pptx → PDF → 逐页 PNG（投影出来的样子）+ 缩略观测表。
* 没有 → **不做视觉验收**，只输出每页的文本框与字号转录，并明确标注"非视觉验收"。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from nibi_kit import cli, contact, engines  # noqa: E402
from nibi_kit.env import require_modules  # noqa: E402

require_modules(["pptx", "PIL"], "pptx_render（python-pptx + pillow）")

EMU_PER_IN = 914400


def render_with_engine(path: Path, outdir: Path, dpi: int) -> dict:
    pdf = engines.convert_to_pdf(path, outdir)
    pages = engines.pdf_to_png(pdf, outdir / "pages", dpi=dpi)
    sheet = contact.make_contact_sheet(pages, outdir / "contact-sheet.png", columns=3,
                                       thumb_width=520)
    return {"mode": "engine", "pdf": str(pdf), "pages": len(pages),
            "files": [str(item) for item in pages], "contactSheet": str(sheet),
            "note": f"真渲染（LibreOffice）：{len(pages)} 页"}


def render_transcript(path: Path, outdir: Path) -> dict:
    from pptx import Presentation  # noqa: PLC0415

    deck = Presentation(str(path))
    outdir.mkdir(parents=True, exist_ok=True)
    target = outdir / "transcript.txt"
    lines = ["# 每页文本与字号转录（**非视觉验收**：本机没有 LibreOffice，版式没验）", ""]
    for index, slide in enumerate(deck.slides, start=1):
        lines.append(f"--- 第 {index} 页 ---")
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            sizes = sorted({run.font.size.pt for paragraph in shape.text_frame.paragraphs
                            for run in paragraph.runs if run.font.size})
            size_text = "/".join(f"{size:g}" for size in sizes) or "?"
            geometry = (f"{shape.left / EMU_PER_IN:.2f},{shape.top / EMU_PER_IN:.2f} "
                        f"{shape.width / EMU_PER_IN:.2f}×{shape.height / EMU_PER_IN:.2f}in")
            lines.append(f"  [{geometry} {size_text}pt] {shape.text_frame.text.strip()[:120]}")
    target.write_text("\n".join(lines), encoding="utf-8")
    return {"mode": "transcript", "pages": 0, "files": [str(target)], "contactSheet": "",
            "note": "非视觉验收：本机没有 LibreOffice，只给出每页文本与字号转录"}


def main() -> int:
    cli.bootstrap()
    parser = argparse.ArgumentParser(description="pptx 逐页渲染（或转录降级）")
    parser.add_argument("--pptx", required=True)
    parser.add_argument("--outdir", required=True)
    parser.add_argument("--dpi", type=int, default=110)
    cli.add_common(parser)
    args = parser.parse_args()

    path = Path(args.pptx).expanduser()
    if not path.is_file():
        print(f"找不到文件：{path}", file=sys.stderr)
        return 2
    outdir = Path(args.outdir).expanduser()
    try:
        result = render_with_engine(path, outdir, args.dpi) \
            if engines.engine_report().can_convert else render_transcript(path, outdir)
    except engines.ConvertError as error:
        print(f"渲染失败：{error}", file=sys.stderr)
        return 1
    lines = [result["note"], f"  产出：{outdir}"]
    if result["contactSheet"]:
        lines.append(f"  观测表：{result['contactSheet']}")
    cli.emit({"summary": "\n".join(lines), **result}, args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
