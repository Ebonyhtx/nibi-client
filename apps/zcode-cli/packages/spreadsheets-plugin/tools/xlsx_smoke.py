"""质量门冒烟：正向对照 + 反向对照，防止判据"静默空转"。

    python tools/xlsx_smoke.py [--keep]

做法：
1. **正向对照**：用 xlsx_build 造一份正常的表 → 质量门必须通过（比了 N 项 N>0）；
2. **反向对照**：直接用 openpyxl 造一份**逐条犯规**的表（窄列出 ###、除以零、伪日期文本、
   本机路径、6pt 字号、打印区域裁掉内容）→ 质量门必须**失败**，且每个期望的判据码都要出现。

反向对照是这脚本的重点：门禁"全绿"有两种可能 —— 东西真好，或者判据压根没在跑。
只有"坏文件必被抓到"能区分这两者。缺任何一条期望判据码 → 退出码非零。
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from nibi_kit import cli  # noqa: E402
from nibi_kit.env import require_modules  # noqa: E402

require_modules(["openpyxl", "PIL"], "xlsx_smoke（openpyxl + pillow）")

PLUGIN_ROOT = Path(__file__).resolve().parent.parent
# 自测夹具**不随插件发布**（用户不该在安装目录里看到"样例报告"）：
# 它们只在仓库里（templates/_dev/），用于开发与打包前验收。
DEV_FIXTURES = PLUGIN_ROOT / "templates" / "_dev"
TOOLS = Path(__file__).resolve().parent

# 反向对照里每个"犯规格"对应的期望判据码。
EXPECTED_CODES = {
    "number-truncated",       # 窄列里的数字 → ####
    "text-clipped",           # 窄列文字被右邻内容裁掉
    "font-too-small",         # 6pt 字号
    "date-as-text",           # 日期写成文本
    "ratio-not-percent",      # 比率列裸小数
    "win-drive-path",         # 单元格里带本机路径（判据码来自 nibi_kit.hygiene 的路径模式表）
    "print-area-clip",        # 打印区域裁掉内容
    "formula-div-zero",       # 除以字面量 0
    "empty-sheet",            # 空工作表
}


def load_report(path: Path) -> dict:
    """读判据报告：`--json-out` 写的是 cli 的 payload（`{"summary", "report"}`），兼容裸报告。"""
    data = json.loads(path.read_text(encoding="utf-8"))
    return data.get("report", data)


def run_tool(script: str, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(TOOLS / script), *args],  # noqa: S603
                          capture_output=True, text=True, encoding="utf-8", errors="replace")


def positive_control(workdir: Path) -> tuple[str, str]:
    spec = DEV_FIXTURES / "spec-quarter.json"
    if not spec.is_file():
        return "skip", ("正向对照需要仓库内的样板（templates/_dev/spec-quarter.json）；"
                        "该夹具不随插件发布，安装目录里跑本项会跳过。")
    target = workdir / "正向对照.xlsx"
    built = run_tool("xlsx_build.py", "--spec", str(spec), "--out", str(target))
    if built.returncode != 0 or not target.is_file():
        return "fail", f"构建失败：{built.stderr.strip()[-300:]}"
    report_path = workdir / "正向对照.判据.json"
    gated = run_tool("xlsx_gate.py", "--xlsx", str(target), "--no-engine",
                     "--json-out", str(report_path))
    if not report_path.is_file():
        return "fail", "质量门没有产出报告"
    report = load_report(report_path)
    ran = [gate for gate in report["gates"] if gate["checked"] > 0]
    if gated.returncode != 0:
        bad = [item for gate in report["gates"] for item in gate["findings"]
               if item["severity"] == "fail"]
        return "fail", "正常文件被拦下：" + json.dumps(bad[:4], ensure_ascii=False)
    if len(ran) < 7:
        return "fail", f"正向对照里只有 {len(ran)} 条判据真的比过东西（应有 ≥7）"
    return "ok", f"{len(ran)} 条判据共比了 {sum(gate['checked'] for gate in ran)} 项"


def make_broken(workdir: Path) -> Path:
    """逐条犯规格的对照文件：每条犯规都对应 EXPECTED_CODES 里的一个判据码。

    表头要写成语义正确的名字（如「日期」、`比率`）——判据是按表头语义判断的，
    表头乱写会让判据"本来就不该管"，而不是"抓不到"。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    target = workdir / "反向对照.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Sheet1"
    ws["A1"] = "日期"
    ws["B1"] = "比率"
    ws["C1"] = "窄列"
    ws["A2"] = "2026-01-05"          # 日期写成文本
    ws["A2"].number_format = "General"
    ws["B2"] = 0.35                  # 比率列裸小数（没设百分比）
    ws["C2"] = 1234567.89            # 窄列里的数字 → ####
    ws["C3"] = "这一格文字很长而右邻有内容于是会被裁掉"
    ws["D3"] = "占位"
    ws["A4"] = "C:\\Users\\<用户名>\\Documents\\数据.xlsx"   # 本机路径
    ws["A5"] = "=1/0"                # 除以字面量 0
    ws["A6"] = 6
    ws["A6"].font = Font(size=6)     # 6pt 字号
    for column, width in (("A", 8), ("B", 8), ("C", 6), ("D", 8)):
        ws.column_dimensions[column].width = width
    ws.print_area = "A1:C3"          # 故意漏掉第 4 行以下的"内容"
    workbook.create_sheet(title="空表")
    workbook.save(target)
    return target


def negative_control(workdir: Path) -> tuple[str, str]:
    target = make_broken(workdir)
    report_path = workdir / "反向对照.判据.json"
    gated = run_tool("xlsx_gate.py", "--xlsx", str(target), "--no-engine",
                     "--json-out", str(report_path))
    if not report_path.is_file():
        return "fail", "质量门没有产出报告"
    report = load_report(report_path)
    codes = {item["code"] for gate in report["gates"] for item in gate["findings"]}
    missing = sorted(EXPECTED_CODES - codes)
    if gated.returncode == 0:
        return "fail", "犯规文件竟然通过了质量门（判据空转）"
    if missing:
        return "fail", "以下判据没有抓到对应的犯规：\n    " + "\n    ".join(
            f"{code}（期望它命中）" for code in missing)
    return "ok", f"{len(EXPECTED_CODES)} 条期望判据全部命中，另有 {len(codes - EXPECTED_CODES)} 条附带命中"


def main() -> int:
    cli.bootstrap()
    parser = argparse.ArgumentParser(description="质量门冒烟（正向 + 反向对照）")
    parser.add_argument("--keep", action="store_true", help="保留中间文件")
    cli.add_common(parser)
    args = parser.parse_args()

    with tempfile.TemporaryDirectory(prefix="nibi-xlsx-smoke-") as raw:
        workdir = Path(raw)
        results: list[tuple[str, str, str]] = []
        status, detail = positive_control(workdir)
        results.append(("正向对照（正常表必须过）", status, detail))
        status, detail = negative_control(workdir)
        results.append(("反向对照（犯规表必须被抓）", status, detail))

        lines = ["xlsx 质量门冒烟："]
        for title, state, detail in results:
            mark = {"ok": "ok  ", "fail": "FAIL", "skip": "skip"}[state]
            lines.append(f"  [{mark}] {title}")
            for chunk in detail.splitlines():
                lines.append(f"         {chunk}")
        failed = [title for title, state, _ in results if state == "fail"]
        skipped = [title for title, state, _ in results if state == "skip"]
        # 跳过必须显式计数：正反向都跳过等于什么都没验，不能算"通过"。
        tail_note = f"（跳过 {len(skipped)} 项：{'、'.join(skipped)}）" if skipped else ""
        lines.append(f"  结论：{'通过' if not failed else '未通过——' + '、'.join(failed)}{tail_note}")
        if args.keep and failed:
            keep = Path.cwd() / "smoke-failure"
            keep.mkdir(exist_ok=True)
            for item in workdir.iterdir():
                (keep / item.name).write_bytes(item.read_bytes())
            lines.append(f"  （失败现场已保留到 {keep}）")
        cli.emit({"summary": "\n".join(lines),
                  "results": [{"case": title, "status": state, "detail": detail}
                              for title, state, detail in results]}, args)
        return 0 if not failed else 1


if __name__ == "__main__":
    raise SystemExit(main())
