"""质量门：交付前把 .xlsx 逐条验一遍，验不到的要显式标「跳过」而不是沉默放过。

## 判据的分层

**A 层（静态，任何环境都能跑）**：直接在文件结构上判
  1. `font-size` 字号下限（含印出等效字号：按「缩放适应一页宽」反推缩放系数）
  2. `overflow` 截断风险（数字溢出成 `###`、文本被非空邻格裁掉、自动换行但行高不够）
  3. `formula` 公式静态检查（括号/引号配对、引用不存在的工作表、`#REF!`、除零）
  4. `number-format` 语义格式（占比列没设百分比、日期是文本、浮点残渣）
  5. `print` 打印就绪（打印区域包含全部内容、重复标题行、页脚页码）
  6. `hygiene` 交付卫生（本机路径/用户名/凭据、核心属性身份）
  7. `structure` 结构卫生（空表、默认表名、used range 虚胖、合并重叠、空表头）

**B 层（引擎，装了 LibreOffice 才跑）**：真渲染
  8. `recalc` 用 LibreOffice 重算 → 读回缓存值，抓真错误值（`#DIV/0!` 等）
  9. `render` 转 PDF → 抽文本抓 `###` 截断与错误 token，并把逐页 PNG 落到证据目录

B 层缺失时报告里是 `skip` + 原因（绝不写成 pass）；`--require-engine` 可把它升级成硬失败。
每条判据都自报「比了多少项」——`checked == 0` 且没写跳过原因会被框架直接判失败。
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from nibi_kit import cli, engines, gates, hygiene, textmetrics  # noqa: E402
from nibi_kit.env import require_modules  # noqa: E402

require_modules(["openpyxl", "PIL"], "xlsx_gate（openpyxl + pillow）")

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter, range_boundaries  # noqa: E402

EXPECTED_GATES = ["font-size", "overflow", "formula", "number-format", "print",
                  "hygiene", "structure", "recalc", "render"]

ERROR_VALUES = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!", "#SPILL!",
                "#CALC!", "#GETTING_DATA")
RATIO_HEADERS = ("率", "比", "占比", "百分比", "份额", "share", "rate", "ratio", "margin", "增速")
MONEY_HEADERS = ("金额", "收入", "成本", "费用", "利润", "销售额", "营收", "price", "amount",
                 "revenue", "cost", "profit", "budget", "预算")
DATE_HEADERS = ("日期", "时间", "date", "time", "月份", "年月")
DATE_TEXT = re.compile(r"^\d{4}[-/年]\d{1,2}([-/月]\d{1,2}日?)?$")
DYNAMIC_FUNCTIONS = ("FILTER(", "UNIQUE(", "XLOOKUP(", "SORT(", "SORTBY(", "SEQUENCE(", "LET(",
                     "LAMBDA(", "RANDARRAY(", "TEXTSPLIT(", "VSTACK(", "HSTACK(")
FONT_FLOOR_HARD = 7.0
FONT_FLOOR_SOFT = 9.0

PAGE_WIDTH_PT = {"portrait": 595.28, "landscape": 841.89}
PAGE_HEIGHT_PT = {"portrait": 841.89, "landscape": 595.28}


# ---------------------------------------------------------------- 工具


def display_preview(value, number_format: str) -> str:
    """按数字格式估算「用户看到的字」（判宽度用；不追求与 Excel 完全一致，宁可偏长）。"""
    if value is None or value == "":
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        # 日期列必须按**实际格式**展开：`yyyy-mm` 显示 7 个字符，不是 10 个 ——
        # 用固定占位串（0000-00-00）会把月列表判成"文字被裁"（假阳性，实测踩过）。
        fmt = number_format or "yyyy-mm-dd"
        if "h" in fmt:
            return f"{value:%Y-%m-%d %H:%M}"
        if "d" in fmt:
            return f"{value:%Y-%m-%d}"
        if "m" in fmt:
            return f"{value:%Y-%m}"
        return f"{value:%Y}"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, str):
        return value
    if not isinstance(value, (int, float)):
        return str(value)
    fmt = number_format or "General"
    if "E+" in fmt or "E-" in fmt:
        return f"{value:.2E}"
    # 多段格式（正数;负数;零）与引号字面量（+0.0"pp";-0.0"pp"）：按段取用，
    # 否则 "+1.2pp" 会被渲染成 "1.2"，读者看到的数与作者意图不同（自渲染实测踩过）。
    sections = fmt.split(";")
    negative = value < 0
    section = sections[1] if (negative and len(sections) > 1) else sections[0]
    if section == "" and negative:
        return ""
    literals = "".join(re.findall(r'"([^"]*)"', section))
    explicit_plus = section.lstrip().startswith("+")
    percent = "%" in section
    match = re.search(r"\.(0+)", section)
    decimals = len(match.group(1)) if match else 0
    scaled = abs(value) * 100 if percent else abs(value)
    text = f"{scaled:,.{decimals}f}" if ("," in section or decimals) else f"{scaled:g}"
    sign = "-" if negative else ("+" if explicit_plus else "")
    prefix = "¥" if "¥" in section else ("$" if "$" in section else "")
    return f"{prefix}{sign}{text}{literals}{'%' if percent else ''}"


def cell_available_px(ws, row: int, column: int, default_width: float) -> tuple[float, int]:
    """一个单元格横向可用像素宽度；被合并时按合并区总宽（返回 (px, 合并区列数)）。"""
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
            total = sum(column_width(ws, index, default_width)
                        for index in range(merged.min_col, merged.max_col + 1))
            return textmetrics.column_width_px(total) - 6.0, merged.max_col - merged.min_col + 1
    return textmetrics.column_width_px(column_width(ws, column, default_width)) - 6.0, 1


def column_width(ws, column: int, default_width: float) -> float:
    letter = get_column_letter(column)
    dimension = ws.column_dimensions.get(letter)
    if dimension is not None and dimension.hidden:
        return 0.0
    if dimension is not None and dimension.width:
        return float(dimension.width)
    return float(default_width)


def content_bounds(ws) -> tuple[int, int, int, int] | None:
    """真实内容范围（不是 openpyxl 的 used range —— 后者会把碰过的空单元格也算进去）。"""
    min_row = min_col = 10**9
    max_row = max_col = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value not in (None, ""):
                min_row = min(min_row, cell.row)
                max_row = max(max_row, cell.row)
                min_col = min(min_col, cell.column)
                max_col = max(max_col, cell.column)
    if max_row == 0:
        return None
    return min_row, min_col, max_row, max_col


# ---------------------------------------------------------------- A 层判据


def gate_font_size(ws_list, report: gates.GateReport, default_width: float) -> None:
    gate = report.gate("font-size", "字号下限与印出等效字号")
    for ws in ws_list:
        fit_scale = print_scale(ws, default_width)
        gate.notes.append(f"{ws.title}：打印缩放 ≈ {fit_scale * 100:.0f}%")
        for row in ws.iter_rows():
            for cell in row:
                if cell.value in (None, ""):
                    continue
                size = float(cell.font.size or 11.0)
                gate.checked += 1
                if size < FONT_FLOOR_HARD:
                    gate.fail("font-too-small", f"字号 {size}pt 小于下限 {FONT_FLOOR_HARD}pt", cell.coordinate)
                elif size < FONT_FLOOR_SOFT:
                    gate.warn("font-small", f"字号 {size}pt 偏小（建议 ≥{FONT_FLOOR_SOFT}pt）", cell.coordinate)
                effective = size * fit_scale
                if effective < FONT_FLOOR_HARD:
                    gate.fail("print-font-too-small",
                              f"{size}pt 经打印缩放（{fit_scale * 100:.0f}%）后等效 {effective:.1f}pt，"
                              f"印出来看不清；请减少列数或分页", cell.coordinate)


def print_scale(ws, default_width: float) -> float:
    """按打印设置反推缩放系数（fitToWidth 时用内容宽度与可用页宽之比）。"""
    setup = ws.page_setup
    if not (ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage):
        scale = float(setup.scale or 100)
        return scale / 100.0
    orientation = (setup.orientation or "landscape").lower()
    usable = PAGE_WIDTH_PT.get(orientation, PAGE_WIDTH_PT["landscape"])
    usable -= 2 * (float(ws.page_margins.left or 0.4) + float(ws.page_margins.right or 0.4)) * 72
    bounds = print_bounds(ws) or content_bounds(ws)
    if bounds is None:
        return 1.0
    _, min_col, _, max_col = bounds
    content_px = sum(column_width(ws, index, default_width) for index in range(min_col, max_col + 1))
    content_pt = textmetrics.column_width_px(content_px) * 0.75
    if content_pt <= usable:
        return 1.0
    fit_width = int(setup.fitToWidth or 1) or 1
    return max(0.1, usable * fit_width / content_pt)


def print_bounds(ws) -> tuple[int, int, int, int] | None:
    area = ws.print_area
    if not area:
        return None
    text = area[0] if isinstance(area, (list, tuple)) else str(area)
    # openpyxl 存回来的打印区域带工作表前缀（`'经营概览'!$A$8:$D$15`），
    # 直接喂 range_boundaries 会 ValueError → 被当成「没设打印区域」（假阴性），先剥前缀。
    if "!" in text:
        text = text.rsplit("!", 1)[1]
    text = text.replace("$", "")
    try:
        # 注意：openpyxl 的 range_boundaries 返回 **(min_col, min_row, max_col, max_row)**，
        # 不是（行,列）—— 直接照行优先解包会把宽表判成"打印区域裁掉了内容"（实测踩过）。
        min_col, min_row, max_col, max_row = range_boundaries(text)
        return min_row, min_col, max_row, max_col
    except ValueError:
        return None


def merged_coordinates(ws) -> set[str]:
    """合并区域内所有单元格坐标。

    openpyxl 的 `merged.cells` 产出的是「行元组」而不是 cell 对象，直接迭代会拿到 tuple
    （`'tuple' object has no attribute 'coordinate'`）—— 这里自己摊平，别踩那个坑。
    """
    coordinates: set[str] = set()
    for merged in ws.merged_cells.ranges:
        for row_index in range(merged.min_row, merged.max_row + 1):
            for column_index in range(merged.min_col, merged.max_col + 1):
                coordinates.add(f"{get_column_letter(column_index)}{row_index}")
    return coordinates


def gate_overflow(ws_list, report: gates.GateReport, default_width: float) -> None:
    gate = report.gate("overflow", "截断与压盖（数字 ### / 文本被裁 / 行高不够）")
    for ws in ws_list:
        merged_cells = merged_coordinates(ws)
        for row in ws.iter_rows():
            for cell in row:
                if cell.value in (None, "") or cell.coordinate in merged_cells:
                    continue
                gate.checked += 1
                available_px, merge_span = cell_available_px(ws, cell.row, cell.column, default_width)
                family = cell.font.name or "微软雅黑"
                size = float(cell.font.size or 11.0)
                bold = bool(cell.font.bold)
                wrap = bool(cell.alignment and cell.alignment.wrap_text)
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    continue
                # 非字符串值一律走 display_preview（按数字格式展开），**不要**用 str()：
                # 日期用 str() 会变成 "2026-01-01 00:00:00"（19 字符），月表被误判成"文字被裁"。
                if not isinstance(cell.value, str):
                    text = display_preview(cell.value, cell.number_format or "")
                    if isinstance(cell.value, (dt.datetime, dt.date)) or (
                            isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)):
                        measured = textmetrics.measure_text(text, family, size, bold=bold)
                        if measured > available_px:
                            gate.fail("number-truncated",
                                      f"数值 {text}（{measured:.0f}px）超过列宽（{available_px:.0f}px），"
                                      f"Excel 会显示 ####", cell.coordinate)
                        continue
                    text = str(cell.value)
                else:
                    text = cell.value
                if "\n" in text and not wrap:
                    gate.warn("newline-without-wrap", "文本含换行但未开自动换行，Excel 只显示第一行",
                              cell.coordinate)
                    text = text.replace("\n", "")
                if wrap:
                    lines = textmetrics.estimate_lines(text, family, size, available_px, bold=bold)
                    if lines > 1:
                        required = textmetrics.required_row_height_pt(lines, size)
                        row_height = ws.row_dimensions[cell.row].height
                        if merge_span > 1:
                            gate.fail("merged-wrap-clipped",
                                      f"合并单元格里的自动换行需要 {lines} 行（约 {required:.0f}pt），"
                                      f"但合并单元格不会自动长高，文字会被压掉", cell.coordinate)
                        elif row_height is not None and float(row_height) < required * 0.95:
                            gate.fail("row-height-too-small",
                                      f"自动换行需 {lines} 行（约 {required:.0f}pt），实际行高 "
                                      f"{float(row_height):.0f}pt，文字被压", cell.coordinate)
                    continue
                measured = textmetrics.measure_text(text, family, size, bold=bold)
                if measured <= available_px:
                    continue
                blocked = _right_neighbour_blocked(ws, cell, merged_cells)
                if blocked is not None:
                    gate.fail("text-clipped",
                              f"文字宽 {measured:.0f}px 超出列宽（{available_px:.0f}px），"
                              f"右侧 {blocked} 有内容 → 会被裁掉", cell.coordinate)
                else:
                    gate.warn("text-overflow",
                              f"文字宽 {measured:.0f}px 超出列宽（{available_px:.0f}px），"
                              f"右侧留空时会溢出显示（不丢字，但会压到邻列观感）", cell.coordinate)


def _right_neighbour_blocked(ws, cell, merged_cells: set[str]) -> str | None:
    """溢出方向上第一个有内容的单元格（Excel 遇非空即裁）。返回坐标或 None。"""
    row = cell.row
    for column in range(cell.column + 1, min(cell.column + 40, ws.max_column + 2)):
        probe = ws.cell(row=row, column=column)
        if probe.value not in (None, ""):
            return probe.coordinate
    return None


def gate_formula(ws_list, report: gates.GateReport) -> None:
    gate = report.gate("formula", "公式静态检查")
    sheet_names = {ws.title.lower() for ws in ws_list}
    for ws in ws_list:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str) or not value.startswith("="):
                    continue
                gate.checked += 1
                formula = value[1:]
                if formula.count("(") != formula.count(")"):
                    gate.fail("formula-unbalanced", f"括号不配对：{value[:60]}", cell.coordinate)
                if formula.count('"') % 2:
                    gate.fail("formula-quote", f"引号不配对：{value[:60]}", cell.coordinate)
                if "#REF!" in formula:
                    gate.fail("formula-ref", f"公式含 #REF!：{value[:60]}", cell.coordinate)
                if re.search(r"/\s*0(?![.\d])", formula):
                    gate.fail("formula-div-zero", f"公式里除以字面量 0：{value[:60]}", cell.coordinate)
                for raw in re.findall(r"(?:'([^']+)'|([A-Za-z0-9_\u4e00-\u9fff]+))!\$?[A-Z]+\$?\d+", formula):
                    name = (raw[0] or raw[1]).lower()
                    if name not in sheet_names:
                        gate.fail("formula-missing-sheet",
                                  f"引用了不存在的工作表 {raw[0] or raw[1]!r}：{value[:60]}", cell.coordinate)
                for function in DYNAMIC_FUNCTIONS:
                    if function in formula.upper():
                        gate.warn("formula-dynamic-array",
                                  f"用了动态数组函数 {function[:-1]}（部分 Excel/WPS 版本不支持）",
                                  cell.coordinate)


def gate_number_format(ws_list, report: gates.GateReport) -> None:
    gate = report.gate("number-format", "数字格式语义")
    for ws in ws_list:
        header_map = _header_names(ws)
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value in (None, "") or (isinstance(value, str) and value.startswith("=")):
                    continue
                gate.checked += 1
                fmt = cell.number_format or "General"
                header = header_map.get(cell.column, "")
                if isinstance(value, float) and not isinstance(value, bool):
                    literal = repr(value)
                    if len(literal.split(".")[-1]) > 8 and "E" not in literal:
                        gate.warn("float-residue",
                                  f"浮点残渣 {literal}（{fmt}）—— 应四舍五入或用 ROUND()", cell.coordinate)
                if isinstance(value, str) and any(key in header for key in DATE_HEADERS) \
                        and DATE_TEXT.match(value.strip()):
                    gate.fail("date-as-text", f"日期值 {value!r} 是文本（无法排序/筛选）—— 应写真日期",
                              cell.coordinate)
                if not isinstance(value, (int, float)) or isinstance(value, bool):
                    continue
                if any(key in header for key in RATIO_HEADERS) and "%" not in fmt and "E" not in fmt \
                        and '"' not in fmt:
                    # 带引号字面量的自定义格式（如 +0.0"pp";-0.0"pp"）是作者有意标注的单位，
                    # 不是"忘了设百分比"；只有裸数字格式才算漏配。
                    if -2 <= float(value) <= 2:
                        gate.fail("ratio-not-percent",
                                  f"列「{header}」是比率语义，但 {cell.coordinate} 格式为 {fmt}"
                                  f"（值 {value} 会显示成 {display_preview(value, fmt)}，"
                                  f"易被读成 {display_preview(value, fmt)}%）—— 应设 pct1/pct2",
                                  cell.coordinate)
                if any(key in header for key in MONEY_HEADERS) and "," not in fmt and "¥" not in fmt \
                        and abs(float(value)) >= 1000:
                    gate.warn("money-format",
                              f"金额列「{header}」的 {cell.coordinate} 未设千分位（{fmt}）", cell.coordinate)


def _header_names(ws) -> dict[int, str]:
    """每列的表头文字：取该列最靠上的加粗/填充行，取不到就用第一行。"""
    names: dict[int, str] = {}
    for column in range(1, min(ws.max_column, 200) + 1):
        chosen = ""
        for row_index in range(1, min(ws.max_row, 30) + 1):
            cell = ws.cell(row=row_index, column=column)
            if cell.value in (None, ""):
                continue
            text = str(cell.value)
            if cell.font and cell.font.bold or (cell.fill and cell.fill.fgColor
                                               and cell.fill.fgColor.rgb not in (None, "00000000")):
                chosen = text
                break
            if not chosen:
                chosen = text
        names[column] = chosen
    return names


def gate_print(ws_list, report: gates.GateReport, default_width: float) -> None:
    gate = report.gate("print", "打印就绪")
    for ws in ws_list:
        bounds = content_bounds(ws)
        if bounds is None:
            continue
        gate.checked += 1
        min_row, min_col, max_row, max_col = bounds
        area = print_bounds(ws)
        if area is not None:
            a_min_row, a_min_col, a_max_row, a_max_col = area
            if (min_row < a_min_row or min_col < a_min_col
                    or max_row > a_max_row or max_col > a_max_col):
                gate.fail("print-area-clip",
                          f"打印区域 {ws.print_area} 未覆盖全部内容"
                          f"（内容到 {get_column_letter(max_col)}{max_row}）—— 导出 PDF 会缺内容",
                          ws.title)
        elif area is None:
            gate.warn("print-area-auto",
                      "未设打印区域（Excel 会按已用范围打印；改过尺寸的旧表很容易多打空页）",
                      ws.title)
        if max_row > 45 or max_col > 9:
            if not ws.print_title_rows:
                gate.warn("no-repeat-header",
                          f"内容 {max_row} 行 × {max_col} 列，跨页打印但未设重复标题行"
                          f"（print.titleRows，如 \"1:2\"）", ws.title)
            if not (ws.oddFooter.center.text if ws.oddFooter and ws.oddFooter.center else ""):
                gate.warn("no-page-number", "跨页打印但页脚没有页码", ws.title)
        if max_col > 9 and (ws.page_setup.orientation or "portrait") == "portrait":
            gate.notes.append(f"{ws.title}：内容 {max_col} 列而纸张是纵向（宽表建议 landscape）")
        gate.checked += _check_charts_in_print_area(ws, area, gate)


def _check_charts_in_print_area(ws, area, gate) -> int:
    """图表整块必须落在打印区域里 —— 锚点在区域外时，导出/打印会**整张图消失**（实测踩过）。"""
    checked = 0
    for chart in getattr(ws, "_charts", []):
        coverage = chart_coverage(ws, chart)
        if coverage is None:
            continue
        checked += 1
        chart_min_row, chart_min_col, chart_max_row, chart_max_col = coverage
        where = (f"{get_column_letter(chart_min_col)}{chart_min_row}"
                 f":{get_column_letter(chart_max_col)}{chart_max_row}")
        if area is None:
            gate.warn("chart-without-print-area",
                      f"图表占位 {where} 而本表没有打印区域（导出时按已用范围走，容易漏图）", ws.title)
            continue
        a_min_row, a_min_col, a_max_row, a_max_col = area
        if (chart_min_row < a_min_row or chart_min_col < a_min_col
                or chart_max_row > a_max_row or chart_max_col > a_max_col):
            gate.fail("chart-outside-print-area",
                      f"图表占位 {where} 超出打印区域 {ws.print_area} → 导出/打印时整张图会消失",
                      ws.title)
    return checked


def chart_coverage(ws, chart) -> tuple[int, int, int, int] | None:
    """图表在表里占的范围（1 基）。支持 TwoCellAnchor（锚点带终点）与 OneCellAnchor（带尺寸）。"""
    anchor = getattr(chart, "anchor", None)
    start = getattr(anchor, "_from", None)
    if start is None:
        return None
    min_row, min_col = int(start.row) + 1, int(start.col) + 1
    end = getattr(anchor, "to", None)
    if end is not None:
        return min_row, min_col, int(end.row) + 1, int(end.col) + 1
    extent = getattr(anchor, "ext", None)
    if extent is None:
        return min_row, min_col, min_row, min_col
    # EMU → px（914400 EMU/英寸，96 px/英寸），再按列宽/行高累加出覆盖范围。
    remaining_w = max(0.0, float(extent.cx)) / 914400 * 96
    default_width = float(ws.sheet_format.defaultColWidth or 8.43)
    column = min_col
    while remaining_w > 0 and column <= 512:
        letter = get_column_letter(column)
        dimension = ws.column_dimensions.get(letter)
        width = float(dimension.width) if (dimension is not None and dimension.width) else default_width
        remaining_w -= textmetrics.column_width_px(width)
        column += 1
    remaining_h = max(0.0, float(extent.cy)) / 914400 * 96
    row = min_row
    while remaining_h > 0 and row <= 50000:
        dimension = ws.row_dimensions.get(row)
        height = float(dimension.height) if (dimension is not None and dimension.height) else 15.0
        remaining_h -= height * 96 / 72
        row += 1
    return min_row, min_col, row - 1, column - 1


def gate_hygiene(ws_list, report: gates.GateReport, workbook, path: Path) -> None:
    gate = report.gate("hygiene", "交付卫生（路径/身份/凭据）")
    parts = hygiene.text_parts(path)
    findings: list[gates.Finding] = []
    for name, text in parts:
        findings += hygiene.scan_office_file_text(text, f"{path.name}:{name}")
    findings += hygiene.core_property_findings(workbook.properties)
    # 比了多少项 = 扫过的 XML 部件数 + 核心属性 1 项。部件数为 0（非 zip / 过滤写错）
    # 会让这条判据变成"空转"，框架会直接把它标红 —— 这是故意留的自曝机制。
    gate.checked = len(parts) + 1
    for finding in findings:
        gate.fail(finding.code, finding.message, finding.location)
    if not findings:
        gate.notes.append(f"扫过 {len(parts)} 个 XML 部件 + 核心属性：没有本机路径/用户名/凭据痕迹")


def gate_structure(ws_list, report: gates.GateReport) -> None:
    gate = report.gate("structure", "结构卫生")
    for ws in ws_list:
        gate.checked += 1
        bounds = content_bounds(ws)
        if bounds is None:
            gate.fail("empty-sheet", "整张工作表没有内容", ws.title)
            continue
        if re.fullmatch(r"Sheet\d*|工作表\d*", ws.title or ""):
            gate.warn("default-sheet-name", f"工作表名是默认名 {ws.title!r}（交付件应改成表意名）", ws.title)
        min_row, min_col, max_row, max_col = bounds
        if ws.max_row and ws.max_row - max_row > 30:
            gate.warn("used-range-fat",
                      f"已用范围到第 {ws.max_row} 行但内容只到第 {max_row} 行 → 尾部空行"
                      f"（打印与体积都会失控）", ws.title)
        if ws.max_column and ws.max_column - max_col > 20:
            gate.warn("used-range-fat-col",
                      f"已用范围到第 {get_column_letter(ws.max_column)} 列但内容只到 "
                      f"{get_column_letter(max_col)} 列 → 尾部空列", ws.title)
        merged = list(ws.merged_cells.ranges)
        for index, first in enumerate(merged):
            for second in merged[index + 1:]:
                if (first.min_row <= second.max_row and first.max_row >= second.min_row
                        and first.min_col <= second.max_col and first.max_col >= second.min_col):
                    gate.fail("merge-overlap", f"合并区域重叠：{first} 与 {second}", ws.title)
        merged_here = merged_coordinates(ws)
        header_row = _find_header_row(ws, min_row, max_row, min_col, max_col, merged_here)
        if header_row is not None:
            blanks = [get_column_letter(index) for index in range(min_col, max_col + 1)
                      if ws.cell(row=header_row, column=index).value in (None, "")
                      and f"{get_column_letter(index)}{header_row}" not in merged_here]
            filled = (max_col - min_col + 1) - len(blanks)
            if max_row > header_row and blanks and filled >= 2:
                gate.warn("blank-header",
                          f"第 {header_row} 行有 {len(blanks)} 个空表头（{'、'.join(blanks[:6])}）", ws.title)


def _find_header_row(ws, min_row: int, max_row: int, min_col: int, max_col: int,
                     merged: set[str]) -> int | None:
    """找表头行：往下最多看 8 行，取第一行「至少有 2 个非合并区里的非空格」的行。

    不能直接用第一行 —— 交付表的第一行常常是跨列标题（合并区），
    拿它当表头会把被合并覆盖的列全报成「空表头」（假阳性）。
    """
    for row_index in range(min_row, min(min_row + 8, max_row) + 1):
        filled = sum(1 for index in range(min_col, max_col + 1)
                     if ws.cell(row=row_index, column=index).value not in (None, "")
                     and f"{get_column_letter(index)}{row_index}" not in merged)
        if filled >= 2:
            return row_index
    return None


# ---------------------------------------------------------------- B 层判据（引擎）


def gate_engine(path: Path, report: gates.GateReport, evidence_dir: Path | None,
                require_engine: bool, dpi: int) -> dict:
    engine = engines.engine_report()
    evidence: dict = {}
    if not engine.can_convert:
        reason = ("本机没有 LibreOffice（soffice），真渲染与公式重算未执行。"
                  "装 LibreOffice 后跑 `python tools/env_doctor.py --office-home <安装目录>` 登记。")
        for gate_id, title in (("recalc", "公式重算（真值）"), ("render", "真渲染（逐页 PDF/PNG）")):
            gate = report.gate(gate_id, title)
            if require_engine:
                gate.add("fail", "engine-missing", reason)
            else:
                gate.skipped = reason
        return evidence

    workdir = (evidence_dir or path.parent / f"{path.stem}.审查") 
    workdir.mkdir(parents=True, exist_ok=True)

    recalc_gate = report.gate("recalc", "公式重算（真值）")
    try:
        recalc_path = engines.recalc_xlsx(path, workdir)
        evidence["recalc"] = str(recalc_path)
        computed = load_workbook(recalc_path, data_only=True)
        try:
            for ws in computed.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value in ERROR_VALUES:
                            recalc_gate.checked += 1
                            recalc_gate.fail("error-value",
                                             f"重算后是错误值 {cell.value}", f"{ws.title}!{cell.coordinate}")
                        elif cell.value not in (None, ""):
                            recalc_gate.checked += 1
        finally:
            computed.close()
        if recalc_gate.checked == 0:
            recalc_gate.notes.append("文件里没有需要求值的单元格（无公式/无数据）")
    except engines.ConvertError as error:
        recalc_gate.add("fail", "recalc-failed", str(error))

    render_gate = report.gate("render", "真渲染（逐页 PDF/PNG）")
    try:
        pdf_path = engines.convert_to_pdf(path, workdir)
        evidence["pdf"] = str(pdf_path)
        page_texts = engines.pdf_text_per_page(pdf_path)
        render_gate.checked += len(page_texts)
        for index, text in enumerate(page_texts, start=1):
            for token in ERROR_VALUES:
                if token in text:
                    render_gate.fail("rendered-error", f"渲染结果里出现错误值 {token}", f"第 {index} 页")
            if re.search(r"#{3,}", text):
                render_gate.fail("rendered-hash", "渲染结果里出现 ###（列宽不足的数字截断）", f"第 {index} 页")
            if "□" in text or "\ufffd" in text:
                render_gate.fail("rendered-missing-glyph", "渲染结果里出现缺字形方框/替换符", f"第 {index} 页")
        png_dir = workdir / "pages"
        pngs = engines.pdf_to_png(pdf_path, png_dir, dpi=dpi)
        evidence["pages"] = [str(item) for item in pngs]
        render_gate.notes.append(f"共 {len(page_texts)} 页，逐页图 {len(pngs)} 张 → {png_dir}")
        if len(page_texts) > 1:
            render_gate.notes.append("多页：请把逐页图交给 visual-judge（或人眼）逐页验收")
    except engines.ConvertError as error:
        render_gate.add("fail", "render-failed", str(error))
    return evidence


# ---------------------------------------------------------------- 入口


def main() -> int:
    cli.bootstrap()
    parser = argparse.ArgumentParser(description="xlsx 质量门")
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--evidence-dir", default="", help="引擎证据（PDF/重算副本/逐页图）落盘目录")
    parser.add_argument("--no-engine", action="store_true", help="只跑 A 层静态判据")
    parser.add_argument("--require-engine", action="store_true", help="引擎缺席即失败")
    parser.add_argument("--dpi", type=int, default=110)
    cli.add_common(parser)
    args = parser.parse_args()

    path = Path(args.xlsx).expanduser()
    if not path.is_file():
        print(f"找不到文件：{path}", file=sys.stderr)
        return 2
    report = gates.GateReport(subject=str(path))
    report.expected(EXPECTED_GATES)
    workbook = load_workbook(path)
    try:
        ws_list = list(workbook.worksheets)
        default_width = float(workbook.worksheets[0].sheet_format.defaultColWidth or 8.43)
        gate_font_size(ws_list, report, default_width)
        gate_overflow(ws_list, report, default_width)
        gate_formula(ws_list, report)
        gate_number_format(ws_list, report)
        gate_print(ws_list, report, default_width)
        gate_hygiene(ws_list, report, workbook, path)
        gate_structure(ws_list, report)
    finally:
        workbook.close()
    evidence: dict = {}
    if not args.no_engine:
        evidence = gate_engine(path, report, Path(args.evidence_dir).expanduser() if args.evidence_dir else None,
                              args.require_engine, args.dpi)
    else:
        for gate_id, title in (("recalc", "公式重算（真值）"), ("render", "真渲染（逐页 PDF/PNG）")):
            report.gate(gate_id, title).skipped = "--no-engine：本次只跑静态判据"

    report.meta["evidence"] = evidence
    # 只通过 cli.emit 写一份（它自己处理 --json-out）：两处都写会互相覆盖，
    # 且后写的那份结构不同 —— 下游读报告会 KeyError（实测踩过）。
    payload = {"summary": report.render(), "report": report.to_dict()}
    cli.emit(payload, args)
    return report.exit_code


if __name__ == "__main__":
    raise SystemExit(main())
